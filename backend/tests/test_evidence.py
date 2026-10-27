import pytest
from pathlib import Path
import openpyxl
from fastapi.testclient import TestClient
from main import app
from services.state_manager import state
from services.evidence_service import EvidenceService

client = TestClient(app)

def test_evidence_1_input_dataset_exists():
    """Verify input dataset exists on disk and is not empty."""
    assert state.current_input_path.exists()
    assert state.current_input_path.stat().st_size > 0

def test_evidence_2_storage_info_endpoint():
    """Verify storage info endpoint returns verifiable file metadata."""
    res = client.get("/api/evidence/storage-info")
    assert res.status_code == 200
    data = res.json()
    assert data["exists"] is True
    assert data["filename"] == "sample_atmospheric_data.csv"
    assert "backend/data" in data["relative_path"]
    assert data["row_count"] > 100
    assert len(data["columns"]) >= 4
    assert len(data["sha256_hash"]) == 64  # valid SHA-256 length

def test_evidence_3_excel_generation_and_all_sheets_exist():
    """Verify atmospheric_monitoring_evidence.xlsx is generated with all 6 required sheets."""
    res = client.get("/api/evidence/excel")
    assert res.status_code == 200
    assert "spreadsheetml.sheet" in res.headers["content-type"]

    excel_path = state.evidence.excel_path
    assert excel_path.exists()
    assert excel_path.stat().st_size > 0

    # Load workbook using openpyxl and verify sheets
    wb = openpyxl.load_workbook(str(excel_path), data_only=True)
    required_sheets = [
        "Raw_Data",
        "Processed_Data",
        "Anomalies",
        "Data_Summary",
        "Validation_Rules",
        "Processing_Log"
    ]
    for sheet_name in required_sheets:
        assert sheet_name in wb.sheetnames, f"Missing required sheet: {sheet_name}"

    # Verify Raw_Data has rows
    ws_raw = wb["Raw_Data"]
    assert ws_raw.max_row > 100

    # Verify Processed_Data row count matches application output
    ws_proc = wb["Processed_Data"]
    assert ws_proc.max_row - 1 == len(state.processor.processed_readings)

    # Verify Anomalies count matches detection logic
    ws_anom = wb["Anomalies"]
    assert ws_anom.max_row - 1 == len(state.processor.anomalies)

def test_evidence_4_file_hashes_and_verification():
    """Verify SHA-256 calculation and PASS/FAIL integrity verification."""
    # List files
    res = client.get("/api/evidence/files")
    assert res.status_code == 200
    files = res.json()
    assert len(files) >= 4  # raw csv, processed csv, anomalies csv, excel, log

    # Verify hash endpoint with valid file
    res_verify = client.post("/api/evidence/verify-hash", json={"filename": "atmospheric_monitoring_evidence.xlsx"})
    assert res_verify.status_code == 200
    verify_data = res_verify.json()
    assert verify_data["status"] == "PASS"
    assert verify_data["match"] is True
    assert len(verify_data["calculated_hash"]) == 64

    # Verify hash with non-existent file
    res_fail = client.post("/api/evidence/verify-hash", json={"filename": "non_existent_file.xyz"})
    assert res_fail.status_code == 200
    fail_data = res_fail.json()
    assert fail_data["status"] == "FAIL"
    assert fail_data["match"] is False

def test_evidence_5_processing_workflow_history():
    """Verify processing history records workflow stages."""
    res = client.get("/api/evidence/history")
    assert res.status_code == 200
    history = res.json()
    assert len(history) > 0
    latest = history[0]
    assert "input_filename" in latest
    assert "input_sha256" in latest
    assert "records_read" in latest
    assert "anomalies_detected" in latest
    assert latest["status"] == "COMPLETED"

def test_evidence_6_raw_dataset_preview():
    """Verify raw preview endpoint returns column names and rows."""
    res = client.get("/api/evidence/raw-preview?limit=10")
    assert res.status_code == 200
    preview = res.json()
    assert len(preview["columns"]) >= 4
    assert len(preview["rows"]) == 10
    assert preview["total_rows"] > 100

def test_evidence_7_downloads():
    """Verify raw CSV and log downloads are accessible."""
    res_raw = client.get("/api/evidence/raw-csv")
    assert res_raw.status_code == 200
    assert "text/csv" in res_raw.headers["content-type"]

    res_log = client.get("/api/evidence/log")
    assert res_log.status_code == 200
    assert "SCIENTIFIC DATASET PROCESSING & AUDIT LOG" in res_log.text

def test_evidence_8_excel_sheet_data_viewer_endpoint():
    """Verify excel-data endpoint parses and returns all 6 sheets for in-browser Excel Sheet Viewer."""
    res = client.get("/api/evidence/excel-data")
    assert res.status_code == 200
    data = res.json()
    assert data["exists"] is True
    assert "Raw_Data" in data["sheet_names"]
    assert "Processed_Data" in data["sheet_names"]
    assert "Anomalies" in data["sheet_names"]
    assert "Data_Summary" in data["sheet_names"]
    assert "Validation_Rules" in data["sheet_names"]
    assert "Processing_Log" in data["sheet_names"]
    assert len(data["sheets"]["Raw_Data"]["headers"]) >= 4
    assert data["sheets"]["Raw_Data"]["row_count"] > 100
    assert len(data["sheets"]["Processed_Data"]["headers"]) >= 6

